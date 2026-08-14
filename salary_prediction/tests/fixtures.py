"""テスト用のダミー Excel（①②③）を生成する。

実データ（給与情報）はリポジトリに置けないので、SPEC.md に書かれた構造と
表記ゆれのパターンを再現した最小限のブックをテストのたびに作る。
"""

from __future__ import annotations

import datetime
import os

import openpyxl
from openpyxl.styles import Color

# ①のレイアウト
HEADER_ROW = 4
FIRST_DAY_ROW = 5


def make_shift_book(path: str, year: int = 2026, month: int = 8) -> str:
    """①シフト表を作る。

    「向笠」…田中(正)/鈴木(パート)/見出し空欄の非常勤列
    「豊岡南1」「豊岡南2」…応援先。田中が2日に応援に行った実績が入っている
    """
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "向笠"
    ws["A1"] = f"{year}年{month}月 シフト予定"
    ws.cell(row=HEADER_ROW, column=1, value="日付")
    ws.cell(row=HEADER_ROW, column=2, value="曜日")
    ws.cell(row=HEADER_ROW, column=3, value="行事予定等")
    ws.cell(row=HEADER_ROW, column=4, value="田中太郎")
    ws.cell(row=HEADER_ROW, column=5, value="鈴木花子")
    ws.cell(row=HEADER_ROW, column=6, value="非常勤")
    # 6列目は見出しが「非常勤」＝セル内に名前が入るプール列
    # 7列目は見出しが空欄だが勤務が書かれている（全列走査が必要なケース）

    rows = {
        1: {4: "7:30-16:30", 5: "13:00〜18:00", 6: "青島凛さん13:00-1800"},
        2: {3: "AM豊田東2、PM豊岡南1", 4: "豊岡南1", 5: "休"},
        3: {4: "８：００～１７：００", 5: "有休", 7: "芥川8:00～13:00"},
        4: {4: "7:30〜11:00\n15:30〜18:00", 5: "×", 6: "鍵屋7:30〜11:00\n15:30〜18:00"},
        5: {4: "9:00-15:00", 5: "欠勤", 6: "蒔田14:00-16:00"},
        6: {4: "調整中", 5: "お盆休み"},
    }
    for day in range(1, 8):
        row = FIRST_DAY_ROW + day - 1
        ws.cell(row=row, column=1, value=datetime.datetime(year, month, day))
        ws.cell(row=row, column=2, value="月火水木金土日"[(day - 1) % 7])
        for col, value in rows.get(day, {}).items():
            ws.cell(row=row, column=col, value=value)

    for sheet_name, person, times in (("豊岡南①", "田中太郎", "9:00-15:30"),
                                      ("豊岡南②", "大村健", "10:00-16:00"),
                                      ("豊岡北", "芥川優", "8:00-13:00")):
        ws2 = wb.create_sheet(sheet_name)
        ws2.cell(row=HEADER_ROW, column=1, value="日付")
        ws2.cell(row=HEADER_ROW, column=4, value=person)
        for day in range(1, 8):
            row = FIRST_DAY_ROW + day - 1
            ws2.cell(row=row, column=1, value=datetime.datetime(year, month, day))
            if day == 2:
                ws2.cell(row=row, column=4, value=times)

    wb.save(path)
    return path


def _make_person_sheet(ws, filled_days: dict[int, tuple[str, str]] | None = None) -> None:
    """③の個人シート（13行目=1日目）。filled_days は事務が入力済みの日。"""
    ws.cell(row=12, column=3, value="出勤")
    ws.cell(row=12, column=5, value="退勤")
    ws.cell(row=12, column=11, value="休憩")
    ws.cell(row=12, column=20, value="特記")
    for day in range(1, 32):
        ws.cell(row=12 + day, column=1, value=day)
    for day, (start, end) in (filled_days or {}).items():
        ws.cell(row=12 + day, column=3, value=start)
        ws.cell(row=12 + day, column=5, value=end)


def _make_salary_sheet(ws, blocks: list[dict]) -> None:
    """③の「給与」シート。B列に「名前」が現れる行から1ブロック。

    blocks の各要素: {"name":..., "基本給":..., "処遇改善":..., "非税通勤":...,
                      "課税通勤":..., "正社員": bool, ...}
    """
    row = 8
    for block in blocks:
        ws.cell(row=row, column=2, value="名前")
        ws.cell(row=row, column=3, value=block.get("name"))
        i = row + 1  # VBA が i+1 してから相対参照するのに合わせる
        ws.cell(row=i + 6, column=2, value=block.get("基本給", 0))
        ws.cell(row=i + 8, column=7, value=block.get("処遇改善", 0))
        ws.cell(row=i + 10, column=2, value=block.get("非税通勤", 0))
        ws.cell(row=i + 10, column=3, value=block.get("課税通勤", 0))
        if block.get("正社員"):
            ws.cell(row=i + 5, column=5, value="法定内残")
            ws.cell(row=i + 6, column=5, value=block.get("法定内残", 0))
            ws.cell(row=i + 6, column=9, value=block.get("業務手当", 0))
            ws.cell(row=i + 8, column=8, value=block.get("残業手当", 0))
            ws.cell(row=i + 10, column=6, value=block.get("不労控除", 0))
            ws.cell(row=i + 6, column=11, value=block.get("休業手当", 0))
        row += 14


def make_kintai_book(
    path: str,
    persons: list[str],
    red_tabs: list[str] | None = None,
    filled: dict[str, dict[int, tuple[str, str]]] | None = None,
    salary_blocks: list[dict] | None = None,
) -> str:
    """③（学童ごとの給与明細シート）を作る。"""
    wb = openpyxl.Workbook()
    wb.active.title = "★常勤"
    wb.create_sheet("☆非常勤")
    wb.create_sheet("一覧")
    _make_salary_sheet(wb.create_sheet("給与"), salary_blocks or [])
    wb.create_sheet("未入力テンプレ")

    for person in persons:
        ws = wb.create_sheet(person)
        _make_person_sheet(ws, (filled or {}).get(person.lstrip("◎")))
        if red_tabs and person in red_tabs:
            ws.sheet_properties.tabColor = Color(rgb="FFFF0000")
    wb.save(path)
    return path


def make_summary_book(path: str, gakudo_names: list[str], start_row: int = 73) -> str:
    """②全社集計（xlsm の代わりに xlsx）。A列に何か、D列に学童名。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "磐田"
    ws.cell(row=start_row - 1, column=1, value="事業所")
    for offset, name in enumerate(gakudo_names):
        row = start_row + offset
        ws.cell(row=row, column=1, value=offset + 1)
        ws.cell(row=row, column=4, value=name)
    wb.create_sheet("浜松")
    wb.save(path)
    return path


def build_all(tmpdir: str) -> dict:
    """テスト一式（①・③×2・②）を tmpdir に作り、パスを返す。"""
    kintai_dir = os.path.join(tmpdir, "kintai")
    os.makedirs(kintai_dir, exist_ok=True)

    shift = make_shift_book(os.path.join(tmpdir, "shift.xlsx"))
    mukasa = make_kintai_book(
        os.path.join(kintai_dir, "__向笠_給与明細シート.xlsx"),
        persons=["◎田中太郎", "鈴木花子", "青島凛", "鍵屋みどり", "蒔田翠"],
        filled={"田中太郎": {1: ("7:30", "16:30")}},
        salary_blocks=[
            {"name": "田中太郎", "基本給": 200000, "処遇改善": 15000, "非税通勤": 4000,
             "課税通勤": 1000, "正社員": True, "法定内残": 3000, "業務手当": 5000,
             "残業手当": 2000, "不労控除": 1000, "休業手当": 0},
            {"name": "鈴木花子", "基本給": 90000, "処遇改善": 5000, "非税通勤": 2000, "課税通勤": 0},
        ],
    )
    toyooka = make_kintai_book(
        os.path.join(kintai_dir, "__豊岡北小_給与明細シート.xlsx"),
        persons=["芥川優", "田中太郎", "大村健"],
        red_tabs=["田中太郎"],  # 向笠と重複するので赤タブ（使っていないシート）
        salary_blocks=[
            {"name": "芥川優", "基本給": 120000, "処遇改善": 8000, "非税通勤": 3000, "課税通勤": 500},
        ],
    )
    summary = make_summary_book(os.path.join(tmpdir, "summary.xlsx"), ["向笠", "豊岡北小", "岩田小"])
    return {"shift": shift, "kintai_dir": kintai_dir, "mukasa": mukasa,
            "toyooka": toyooka, "summary": summary}
