"""ダミーの①②③を生成して、パース〜書き込み〜集計まで一通り検証する。

    python -m unittest discover -s tests   （salary_prediction フォルダで実行）
"""

from __future__ import annotations

import datetime
import json
import os
import sys
import tempfile
import unittest
import unittest.mock

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from kyuyo import cli, fixed, jp_holidays, textutil as tu, wizard  # noqa: E402
from kyuyo.kintai_index import KintaiIndex  # noqa: E402
from kyuyo.kintai_write import apply_plans  # noqa: E402
from kyuyo.plan import build_plans, compute_break_minutes  # noqa: E402
from kyuyo.shift import DayEntry, ShiftBook, merge_entries  # noqa: E402
from kyuyo.totals import compute_gakudo_totals, write_to_summary  # noqa: E402
from tests import fixtures  # noqa: E402

T = datetime.time


class TestTextUtil(unittest.TestCase):
    def test_time_range_variants(self):
        cases = {
            "7:30-16:30": [(T(7, 30), T(16, 30))],
            "13:00〜18:00": [(T(13, 0), T(18, 0))],
            "8:00～13:00": [(T(8, 0), T(13, 0))],
            "８：００～１３：００": [(T(8, 0), T(13, 0))],
            "13:00-1800": [(T(13, 0), T(18, 0))],
            "830-1230": [(T(8, 30), T(12, 30))],
            "13-18": [(T(13, 0), T(18, 0))],
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(tu.parse_time_ranges(text), expected)

    def test_split_shift(self):
        got = tu.parse_time_ranges("鍵屋7:30〜11:00\n15:30〜18:00")
        self.assertEqual(got, [(T(7, 30), T(11, 0)), (T(15, 30), T(18, 0))])

    def test_invalid_time_is_ignored(self):
        self.assertEqual(tu.parse_time_ranges("25:00-99:99"), [])
        self.assertEqual(tu.parse_time_ranges("調整中"), [])

    def test_absence_and_holiday(self):
        self.assertEqual(tu.match_absence("有給"), "有休")
        self.assertEqual(tu.match_absence("欠勤"), "欠勤")
        self.assertEqual(tu.match_absence("休出"), "休出")
        self.assertIsNone(tu.match_absence("休"))
        self.assertTrue(tu.is_holiday_mark("休"))
        self.assertTrue(tu.is_holiday_mark("×"))
        self.assertTrue(tu.is_holiday_mark("お盆休み"))
        self.assertFalse(tu.is_holiday_mark("有休"))

    def test_gakudo_normalize(self):
        self.assertEqual(tu.normalize_gakudo("豊岡南①"), ("豊岡南", 1))
        self.assertEqual(tu.normalize_gakudo("豊岡南1"), ("豊岡南", 1))
        self.assertEqual(tu.normalize_gakudo("岩田小"), ("岩田", None))
        self.assertEqual(tu.gakudo_key("豊田北部③"), "豊田北部")

    def test_person_normalize(self):
        self.assertEqual(tu.normalize_person("◎田中太郎"), "田中太郎")
        self.assertEqual(tu.normalize_person("青島凛さん"), "青島凛")
        self.assertEqual(tu.normalize_person(" 鈴木 花子 "), "鈴木花子")


class TestBreakRule(unittest.TestCase):
    def test_over_six_hours(self):
        self.assertEqual(compute_break_minutes([(T(8, 0), T(17, 0))]), 60)

    def test_six_hours_or_less(self):
        self.assertEqual(compute_break_minutes([(T(13, 0), T(18, 0))]), 0)
        self.assertEqual(compute_break_minutes([(T(9, 0), T(15, 0))]), 0)

    def test_split_shift_gap_counts_as_break(self):
        # 7:30-11:00(3.5h) + 15:30-18:00(2.5h) = 勤務6h、間の4.5h(270分)は休憩扱い
        self.assertEqual(compute_break_minutes([(T(7, 30), T(11, 0)), (T(15, 30), T(18, 0))]), 270)


class PipelineTestCase(unittest.TestCase):
    """①→③→② を通しで検証するための共通セットアップ。"""

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.paths = fixtures.build_all(cls._tmp.name)
        cls.book = ShiftBook(cls.paths["shift"])
        entries = cls.book.parse_all(["向笠"])
        cls.book.resolve_transfers(entries)
        cls.merged = merge_entries(entries)
        cls.index = KintaiIndex.from_dir(cls.paths["kintai_dir"])
        cls.plans = build_plans(cls.merged, cls.index)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def entry(self, person: str, day: int):
        return self.merged[(person, day)]

    def plan_for(self, person: str, day: int):
        for plan in self.plans:
            if plan.person == person and plan.day == day:
                return plan
        raise AssertionError(f"{person} {day}日の計画が無い")


class TestShiftParsing(PipelineTestCase):
    def test_plain_time_range(self):
        entry = self.entry("田中太郎", 1)
        self.assertEqual(entry.kind, "time")
        self.assertEqual(entry.shifts, [(T(7, 30), T(16, 30))])

    def test_fullwidth_time_range(self):
        self.assertEqual(self.entry("田中太郎", 3).shifts, [(T(8, 0), T(17, 0))])

    def test_split_shift_entry(self):
        entry = self.entry("田中太郎", 4)
        self.assertEqual(entry.shifts, [(T(7, 30), T(11, 0)), (T(15, 30), T(18, 0))])

    def test_absence_goes_to_note_column(self):
        self.assertEqual(self.entry("鈴木花子", 3).kind, "absence")
        self.assertEqual(self.entry("鈴木花子", 3).absence, "有休")
        self.assertEqual(self.entry("鈴木花子", 5).absence, "欠勤")

    def test_absence_cell_is_not_flagged_as_embedded_name(self):
        # 「有休」「欠勤」等のキーワードを氏名と誤認して警告を出さないこと
        self.assertEqual(self.entry("鈴木花子", 3).warnings, [])
        self.assertEqual(self.plan_for("鈴木花子", 3).status, "write")
        self.assertEqual(self.plan_for("鈴木花子", 3).note, "有休")

    def test_holiday_is_not_written(self):
        for day in (2, 4, 6):
            self.assertEqual(self.entry("鈴木花子", day).kind, "holiday")
        self.assertEqual(self.plan_for("鈴木花子", 2).status, "skip")

    def test_name_embedded_in_pool_column(self):
        # 見出し「非常勤」の列に「青島凛さん13:00-1800」
        entry = self.entry("青島凛", 1)
        self.assertEqual(entry.shifts, [(T(13, 0), T(18, 0))])

    def test_headerless_column_is_scanned(self):
        # 見出しが空欄の列に「芥川8:00～13:00」
        self.assertEqual(self.entry("芥川", 3).shifts, [(T(8, 0), T(13, 0))])

    def test_cross_reference_to_other_gakudo(self):
        # 「豊岡南1」とだけ書かれた日は、豊岡南①シートの実際の時刻を引く
        entry = self.entry("田中太郎", 2)
        self.assertEqual(entry.kind, "time")
        self.assertEqual(entry.shifts, [(T(9, 0), T(15, 30))])
        self.assertTrue(any("豊岡南" in note for note in entry.notes))

    def test_uninterpretable_cell_is_flagged(self):
        entry = self.entry("田中太郎", 6)
        self.assertEqual(entry.kind, "unknown")
        self.assertEqual(self.plan_for("田中太郎", 6).status, "review")

    def test_event_column_is_ignored(self):
        # C列の「AM豊田東2、PM豊岡南1」は学童全体の予定なので勤務にしない
        self.assertNotIn(("豊田東", 2), self.merged)


class TestNameIndex(PipelineTestCase):
    def test_red_tab_sheet_is_excluded(self):
        ref, reason = self.index.resolve("田中太郎")
        self.assertIsNotNone(ref)
        self.assertIn("向笠", ref.filename)
        self.assertEqual(ref.sheet, "◎田中太郎")
        self.assertEqual(reason, "氏名完全一致")
        self.assertEqual(self.index.duplicates(), {})
        self.assertEqual([r.sheet for r in self.index.red_tabs()], ["田中太郎"])

    def test_management_sheets_are_excluded(self):
        for name in ("★常勤", "☆非常勤", "一覧", "給与", "未入力テンプレ"):
            self.assertNotIn(name, self.index.by_person)

    def test_partial_match(self):
        ref, reason = self.index.resolve("芥川")
        self.assertEqual(ref.sheet, "芥川優")
        self.assertIn("部分一致", reason)

    def test_unknown_person(self):
        ref, reason = self.index.resolve("存在しない人")
        self.assertIsNone(ref)
        self.assertEqual(reason, "③に該当シートなし")


class TestPlanAndWrite(PipelineTestCase):
    def test_row_mapping(self):
        # 13行目=1日目
        self.assertEqual(self.plan_for("田中太郎", 3).row, 15)

    def test_from_day_filter(self):
        plans = build_plans(self.merged, self.index, from_day=4)
        skipped = [p for p in plans if p.day < 4]
        self.assertTrue(skipped)
        self.assertTrue(all(p.status == "skip" for p in skipped))

    def test_write_creates_copy_and_fills_cells(self):
        out_dir = os.path.join(self._tmp.name, "out_kintai")
        plans = build_plans(self.merged, self.index)
        result = apply_plans(plans, output_dir=out_dir)
        self.assertGreater(result.written, 0)

        path = os.path.join(out_dir, os.path.basename(self.paths["mukasa"]))
        wb = openpyxl.load_workbook(path)
        ws = wb["◎田中太郎"]
        self.assertEqual(ws.cell(row=15, column=3).value, T(8, 0))    # 3日 出勤
        self.assertEqual(ws.cell(row=15, column=5).value, T(17, 0))   # 3日 退勤
        self.assertEqual(ws.cell(row=15, column=11).value, 60)        # 9時間 → 休憩60分
        self.assertEqual(ws.cell(row=16, column=11).value, 270)       # 4日 分割シフトの空き時間

        note_ws = wb["鈴木花子"]
        self.assertEqual(note_ws.cell(row=15, column=20).value, "有休")
        self.assertIsNone(note_ws.cell(row=15, column=3).value)
        wb.close()

        # 原本には触れていない
        original = openpyxl.load_workbook(self.paths["mukasa"])
        self.assertIsNone(original["◎田中太郎"].cell(row=15, column=3).value)
        original.close()

    def test_existing_input_is_not_overwritten(self):
        out_dir = os.path.join(self._tmp.name, "out_keep")
        plans = build_plans(self.merged, self.index)
        result = apply_plans(plans, output_dir=out_dir)
        self.assertGreaterEqual(result.skipped_existing, 1)

        path = os.path.join(out_dir, os.path.basename(self.paths["mukasa"]))
        wb = openpyxl.load_workbook(path)
        # 1日は事務が入力済み（文字列 "7:30"）なので、そのまま残っている
        self.assertEqual(wb["◎田中太郎"].cell(row=13, column=3).value, "7:30")
        wb.close()

    def test_overwrite_flag(self):
        out_dir = os.path.join(self._tmp.name, "out_overwrite")
        plans = build_plans(self.merged, self.index)
        apply_plans(plans, output_dir=out_dir, overwrite_filled=True)
        path = os.path.join(out_dir, os.path.basename(self.paths["mukasa"]))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb["◎田中太郎"].cell(row=13, column=3).value, T(7, 30))
        wb.close()


class TestTotals(PipelineTestCase):
    def test_seishain_branch(self):
        totals = compute_gakudo_totals(self.paths["mukasa"])
        # 正社員: 200000+3000+5000+2000-1000+0 = 209000、パート: 90000
        self.assertEqual(totals.aa_salary, 299000)
        self.assertEqual(totals.bb_kaizen, 20000)
        self.assertEqual(totals.cc_transit, 7000)
        self.assertEqual([b["名前"] for b in totals.blocks], ["田中太郎", "鈴木花子"])

    def test_part_time_template(self):
        totals = compute_gakudo_totals(self.paths["toyooka"])
        self.assertEqual(totals.aa_salary, 120000)
        self.assertEqual(totals.cc_transit, 3500)

    def test_write_to_summary(self):
        rows = {
            "向笠": compute_gakudo_totals(self.paths["mukasa"]),
            "豊岡北小": compute_gakudo_totals(self.paths["toyooka"]),
        }
        out = os.path.join(self._tmp.name, "summary_out.xlsx")
        results = write_to_summary(self.paths["summary"], "磐田", rows,
                                   start_row=73, output_path=out)
        self.assertTrue(all("行 に" in message for message in results.values()))

        wb = openpyxl.load_workbook(out)
        ws = wb["磐田"]
        self.assertEqual(ws.cell(row=73, column=5).value, 299000)   # 向笠 E列
        self.assertEqual(ws.cell(row=73, column=6).value, 20000)    # F列
        self.assertEqual(ws.cell(row=73, column=7).value, 7000)     # G列
        self.assertEqual(ws.cell(row=74, column=5).value, 120000)   # 豊岡北小
        self.assertIsNone(ws.cell(row=75, column=5).value)          # 岩田小は対象外
        wb.close()

    def test_dry_run_does_not_write(self):
        rows = {"向笠": compute_gakudo_totals(self.paths["mukasa"])}
        write_to_summary(self.paths["summary"], "磐田", rows, start_row=73, dry_run=True)
        wb = openpyxl.load_workbook(self.paths["summary"])
        self.assertIsNone(wb["磐田"].cell(row=73, column=5).value)
        wb.close()


class TestCliErrors(unittest.TestCase):
    """ファイルの置き忘れ等で、長いトレースバックではなく説明を出すこと。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_missing_shift_file(self):
        with self.assertRaises(SystemExit) as ctx:
            cli.main(["plan", "--shift", os.path.join(self.tmp.name, "shift.xlsx"),
                      "--kintai-dir", self.tmp.name])
        self.assertIn("①シフト表が見つかりません", str(ctx.exception))

    def test_missing_kintai_dir(self):
        shift = fixtures.make_shift_book(os.path.join(self.tmp.name, "shift.xlsx"))
        with self.assertRaises(SystemExit) as ctx:
            cli.main(["plan", "--shift", shift,
                      "--kintai-dir", os.path.join(self.tmp.name, "kintai")])
        self.assertIn("③のフォルダが見つかりません", str(ctx.exception))

    def test_empty_kintai_dir(self):
        shift = fixtures.make_shift_book(os.path.join(self.tmp.name, "shift.xlsx"))
        empty = os.path.join(self.tmp.name, "kintai")
        os.makedirs(empty)
        with self.assertRaises(SystemExit) as ctx:
            cli.main(["plan", "--shift", shift, "--kintai-dir", empty])
        self.assertIn("③ファイル", str(ctx.exception))

    def test_not_an_xlsx(self):
        broken = os.path.join(self.tmp.name, "shift.xlsx")
        with open(broken, "w", encoding="utf-8") as f:
            f.write("これはExcelではない")
        with self.assertRaises(SystemExit) as ctx:
            cli.main(["plan", "--shift", broken, "--kintai-dir", self.tmp.name])
        self.assertIn("開けません", str(ctx.exception))


class TestJapaneseHolidays(unittest.TestCase):
    def test_2026_holidays(self):
        h = jp_holidays.holidays(2026)
        self.assertEqual(h[datetime.date(2026, 1, 1)], "元日")
        self.assertEqual(h[datetime.date(2026, 1, 12)], "成人の日")       # 1月第2月曜
        self.assertEqual(h[datetime.date(2026, 3, 20)], "春分の日")
        self.assertEqual(h[datetime.date(2026, 8, 11)], "山の日")
        self.assertEqual(h[datetime.date(2026, 9, 21)], "敬老の日")
        self.assertEqual(h[datetime.date(2026, 9, 23)], "秋分の日")

    def test_substitute_holiday(self):
        # 2026-05-03(憲法記念日)が日曜 → 5/6が振替休日（5/4,5/5は祝日のため）
        self.assertEqual(jp_holidays.holiday_name(datetime.date(2026, 5, 6)), "振替休日")

    def test_citizens_holiday(self):
        # 敬老の日(9/21)と秋分の日(9/23)に挟まれた9/22
        self.assertEqual(jp_holidays.holiday_name(datetime.date(2026, 9, 22)), "国民の休日")

    def test_plain_weekday_is_not_holiday(self):
        self.assertFalse(jp_holidays.is_holiday(datetime.date(2026, 8, 17)))


class TestFixedShifts(unittest.TestCase):
    """①に出てこない/毎月同じ勤務の先生の設定。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _config(self, data: dict) -> str:
        path = os.path.join(self.tmp.name, "fixed_shifts.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        return path

    def test_parse_weekdays(self):
        self.assertEqual(fixed.parse_weekdays("平日"), {0, 1, 2, 3, 4})
        self.assertEqual(fixed.parse_weekdays("土"), {5})
        self.assertEqual(fixed.parse_weekdays("土日"), {5, 6})
        self.assertEqual(fixed.parse_weekdays("月-金"), {0, 1, 2, 3, 4})
        self.assertEqual(fixed.parse_weekdays(["月", "水"]), {0, 2})
        self.assertEqual(fixed.parse_weekdays("月曜日"), {0})
        with self.assertRaises(fixed.FixedShiftConfigError):
            fixed.parse_weekdays("なんとなく")

    def test_weekday_rule_skips_weekends_and_holidays(self):
        path = self._config({"rules": [{
            "person": "山田花子", "file": "【統括】", "weekdays": "平日",
            "time": "9:30-18:30", "break_minutes": 60,
        }]})
        entries = fixed.expand(fixed.load_config(path), 2026, 8)
        days = sorted(e.day for e in entries)
        self.assertIn(17, days)          # 月曜
        self.assertNotIn(15, days)       # 土曜
        self.assertNotIn(16, days)       # 日曜
        self.assertNotIn(11, days)       # 山の日
        self.assertEqual(entries[0].shifts, [(T(9, 30), T(18, 30))])
        self.assertEqual(entries[0].break_minutes, 60)
        self.assertEqual(entries[0].file_hint, "【統括】")

    def test_saturday_only_rule(self):
        path = self._config({"rules": [{
            "person": "佐藤一郎", "weekdays": "土", "time": "13:00-18:00",
        }]})
        entries = fixed.expand(fixed.load_config(path), 2026, 8)
        saturdays = [d for d in range(1, 32) if datetime.date(2026, 8, d).weekday() == 5]
        self.assertEqual(sorted(e.day for e in entries), saturdays)
        self.assertEqual(entries[0].shifts, [(T(13, 0), T(18, 0))])
        self.assertIsNone(entries[0].break_minutes)  # 5時間なので通常ルール（休憩なし）

    def test_extra_holidays(self):
        path = self._config({
            "extra_holidays": ["2026-08-13", "2026-08-14"],
            "rules": [{"person": "山田花子", "weekdays": "平日", "time": "9:30-18:30"}],
        })
        days = [e.day for e in fixed.expand(fixed.load_config(path), 2026, 8)]
        self.assertNotIn(13, days)
        self.assertNotIn(14, days)
        self.assertIn(17, days)

    def test_shift_sheet_takes_priority(self):
        merged = {("山田花子", 3): DayEntry(day=3, person="山田花子", kind="time",
                                          shifts=[(T(8, 0), T(12, 0))])}
        path = self._config({"rules": [{
            "person": "山田花子", "weekdays": "平日", "time": "9:30-18:30"}]})
        added = fixed.merge_into(merged, fixed.expand(fixed.load_config(path), 2026, 8))
        self.assertEqual(merged[("山田花子", 3)].shifts, [(T(8, 0), T(12, 0))])
        self.assertTrue(any("①の記載を優先" in n for n in merged[("山田花子", 3)].notes))
        self.assertGreater(added, 0)

    def test_broken_config_is_reported(self):
        path = self._config({"rules": [{"person": "山田花子", "weekdays": "平日"}]})
        with self.assertRaises(fixed.FixedShiftConfigError):
            fixed.load_config(path)

    def test_example_config_is_valid(self):
        example = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                               "fixed_shifts.example.json")
        config = fixed.load_config(example)
        self.assertTrue(config.rules)

    def test_written_into_kintai_with_file_hint(self):
        """固定シフトが、指定した③ファイルの本人シートに書き込まれること。"""
        paths = fixtures.build_all(self.tmp.name)
        toukatsu = os.path.join(paths["kintai_dir"], "■【統括】給与明細シート.xlsx")
        fixtures.make_kintai_book(toukatsu, persons=["山田花子"])

        path = self._config({"rules": [{
            "person": "山田花子", "file": "【統括】", "weekdays": "平日",
            "time": "9:30-18:30", "break_minutes": 60,
        }]})
        out_dir = os.path.join(self.tmp.name, "out_fixed")
        code = cli.main(["plan", "--shift", paths["shift"], "--kintai-dir", paths["kintai_dir"],
                         "--fixed", path, "--out-dir", out_dir])
        self.assertEqual(code, 0)

        code = cli.main(["apply", "--shift", paths["shift"], "--kintai-dir", paths["kintai_dir"],
                         "--fixed", path, "--out-dir", out_dir, "--yes"])
        self.assertEqual(code, 0)

        wb = openpyxl.load_workbook(os.path.join(out_dir, "kintai", os.path.basename(toukatsu)))
        ws = wb["山田花子"]
        # ①は2026年8月。8/3(月)は13+3-1=15行目
        self.assertEqual(ws.cell(row=15, column=3).value, T(9, 30))
        self.assertEqual(ws.cell(row=15, column=5).value, T(18, 30))
        self.assertEqual(ws.cell(row=15, column=11).value, 60)
        # 8/11(山の日)は書き込まれない
        self.assertIsNone(ws.cell(row=23, column=3).value)
        wb.close()


class TestWizard(unittest.TestCase):
    """「開始」から呼ばれる対話モード。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = os.path.join(self.tmp.name, "input")

    def test_reports_missing_files(self):
        wizard.ensure_folders(self.base)
        found = wizard.find_inputs(self.base)
        self.assertIsNone(found["shift"])
        self.assertTrue(any("①シフト表がありません" in p for p in found["problems"]))
        self.assertTrue(any("③がありません" in p for p in found["problems"]))

    def test_finds_files_by_name(self):
        wizard.ensure_folders(self.base)
        fixtures.make_shift_book(os.path.join(self.base, "shift.xlsx"))
        fixtures.make_kintai_book(os.path.join(self.base, "kintai", "__向笠_給与明細シート.xlsx"),
                                  persons=["鈴木花子"])
        fixtures.make_summary_book(os.path.join(self.base, "summary.xlsm"), ["向笠"])
        found = wizard.find_inputs(self.base)
        self.assertEqual(found["problems"], [])
        self.assertTrue(found["shift"].endswith("shift.xlsx"))
        self.assertTrue(found["summary"].endswith("summary.xlsm"))
        self.assertEqual(len(found["kintai_files"]), 1)

    def test_ambiguous_shift_file(self):
        wizard.ensure_folders(self.base)
        fixtures.make_shift_book(os.path.join(self.base, "①.xlsx"))
        fixtures.make_shift_book(os.path.join(self.base, "①(修正).xlsx"))
        fixtures.make_kintai_book(os.path.join(self.base, "kintai", "__向笠_給与明細シート.xlsx"),
                                  persons=["鈴木花子"])
        found = wizard.find_inputs(self.base)
        self.assertIsNone(found["shift"])
        self.assertTrue(any("特定できません" in p for p in found["problems"]))

    def test_single_unnamed_shift_file_is_used(self):
        wizard.ensure_folders(self.base)
        fixtures.make_shift_book(os.path.join(self.base, "8月シフト.xlsx"))
        fixtures.make_kintai_book(os.path.join(self.base, "kintai", "__向笠_給与明細シート.xlsx"),
                                  persons=["鈴木花子"])
        found = wizard.find_inputs(self.base)
        self.assertTrue(found["shift"].endswith("8月シフト.xlsx"))

    def test_full_run_writes_files(self):
        """「開始」→ Enter を押していく流れを、入力を差し替えて再現する。"""
        wizard.ensure_folders(self.base)
        paths = fixtures.build_all(self.tmp.name)
        os.replace(paths["shift"], os.path.join(self.base, "shift.xlsx"))
        for name in os.listdir(paths["kintai_dir"]):
            os.replace(os.path.join(paths["kintai_dir"], name),
                       os.path.join(self.base, "kintai", name))
        os.replace(paths["summary"], os.path.join(self.base, "summary.xlsm"))

        answers = iter(["1", "y", "y", "磐田"])  # 中締め翌日 / 書き込む / 保存した / ②のシート名
        out_dir = os.path.join(self.tmp.name, "out")
        with unittest.mock.patch.object(wizard, "ask", lambda *a, **k: next(answers)), \
             unittest.mock.patch.object(wizard, "open_in_explorer", lambda path: None):
            code = wizard.run(self.base, out_dir)

        self.assertEqual(code, 0)
        self.assertTrue(os.path.isfile(os.path.join(out_dir, "plan.csv")))
        self.assertTrue(os.path.isfile(os.path.join(out_dir, "kintai", "__向笠_給与明細シート.xlsx")))
        wb = openpyxl.load_workbook(os.path.join(out_dir, "kintai", "__向笠_給与明細シート.xlsx"))
        self.assertEqual(wb["◎田中太郎"].cell(row=15, column=3).value, T(8, 0))
        wb.close()

    def test_stops_when_user_declines(self):
        wizard.ensure_folders(self.base)
        paths = fixtures.build_all(self.tmp.name)
        os.replace(paths["shift"], os.path.join(self.base, "shift.xlsx"))
        for name in os.listdir(paths["kintai_dir"]):
            os.replace(os.path.join(paths["kintai_dir"], name),
                       os.path.join(self.base, "kintai", name))

        answers = iter(["1", "n"])
        out_dir = os.path.join(self.tmp.name, "out_declined")
        with unittest.mock.patch.object(wizard, "ask", lambda *a, **k: next(answers)), \
             unittest.mock.patch.object(wizard, "open_in_explorer", lambda path: None):
            code = wizard.run(self.base, out_dir)

        self.assertEqual(code, 0)
        self.assertFalse(os.path.exists(os.path.join(out_dir, "kintai")))


if __name__ == "__main__":
    unittest.main()
