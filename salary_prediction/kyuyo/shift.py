"""①シフト表（xlsx）の読み取り。

学童ごとにシートが分かれており、4行目が見出し（先生名 or 空欄 or「非常勤」）、
5行目以降が1日1行。セルの表記ゆれの分類は SPEC.md「①シフト表の読み取りルール」に従う。

分類（DayEntry.kind）:
    "time"      … 時刻レンジが読めた（分割シフトは複数件）
    "absence"   … 有休・欠勤・休出・慶弔・休業 → ③のT列へ
    "holiday"   … 休・×・お盆休み等 → ③には何も書かない
    "transfer"  … 学童名のみ（応援）→ 応援先シートをクロスリファレンスして時刻を探す
    "unknown"   … 解釈できなかった（要確認リストへ）
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field

import openpyxl

from . import textutil as tu

# ①の既定レイアウト（SPEC.md より）
HEADER_ROW = 4
FIRST_DAY_ROW = 5
DAY_COL = 1
WEEKDAY_COL = 2
EVENT_COL = 3          # C列＝行事予定等（個人の勤務ではないのでパース対象外）
FIRST_PERSON_COL = 4   # D列以降が各先生の勤務時間

# 見出しが人名ではないことが分かっている語
NON_PERSON_HEADERS = {"非常勤", "常勤", "アルバイト", "パート", "応援", "備考", "行事", "予定", "日付", "曜日"}


@dataclass
class DayEntry:
    """①の1セルから読み取った、ある先生の1日分の勤務情報。"""

    day: int
    person: str
    kind: str
    shifts: list[tuple[datetime.time, datetime.time]] = field(default_factory=list)
    absence: str | None = None
    destination: str | None = None
    break_minutes: int | None = None    # 固定シフト等で休憩が決まっている場合
    file_hint: str | None = None        # 書き込み先の③ファイル名（部分一致）
    sources: list[str] = field(default_factory=list)   # 読み取り元（レポート用）
    raws: list[str] = field(default_factory=list)      # 元のセル文字列
    notes: list[str] = field(default_factory=list)     # 解釈の根拠
    warnings: list[str] = field(default_factory=list)  # 目視確認してほしい点

    @property
    def worked_minutes(self) -> int:
        return sum(tu.minutes_between(s, e) for s, e in self.shifts)

    @property
    def needs_review(self) -> bool:
        return bool(self.warnings) or self.kind in ("unknown", "transfer")


class ShiftBook:
    """①シフト表のブック。シート＝学童。"""

    def __init__(self, path: str, header_row: int = HEADER_ROW, first_day_row: int = FIRST_DAY_ROW):
        self.path = path
        self.header_row = header_row
        self.first_day_row = first_day_row
        self._wb = openpyxl.load_workbook(path, data_only=True)
        self.sheet_names: list[str] = list(self._wb.sheetnames)
        # 学童の基本名 → その学童のシート名（枝番違いで複数ある）
        self.gakudo_sheets: dict[str, list[str]] = {}
        for name in self.sheet_names:
            self.gakudo_sheets.setdefault(tu.gakudo_key(name), []).append(name)

    # ------------------------------------------------------------ シート内の構造

    def day_rows(self, sheet_name: str) -> dict[int, int]:
        """{日: 行番号}。A列の日付から求める（値が日付でなければ行の並び順を日とする）。"""
        ws = self._wb[sheet_name]
        rows: dict[int, int] = {}
        blank_run = 0
        row = self.first_day_row
        while row <= ws.max_row and blank_run < 3:
            v = ws.cell(row=row, column=DAY_COL).value
            if v is None or str(v).strip() == "":
                blank_run += 1
                row += 1
                continue
            blank_run = 0
            day = None
            if isinstance(v, (datetime.datetime, datetime.date)):
                day = v.day
            else:
                text = tu.normalize(v)
                digits = "".join(ch for ch in text if ch.isdigit())
                if digits:
                    day = int(digits[:2]) if len(digits) > 2 else int(digits)
            if day is None or not (1 <= day <= 31):
                day = row - self.first_day_row + 1
            rows.setdefault(day, row)
            row += 1
        return rows

    def period(self, sheet_names: list[str] | None = None) -> tuple[int, int] | None:
        """①のA列の日付から (年, 月) を推定する。日付が入っていなければ None。"""
        counts: dict[tuple[int, int], int] = {}
        for sheet_name in (sheet_names or self.sheet_names):
            ws = self._wb[sheet_name]
            for row in range(self.first_day_row, min(ws.max_row, self.first_day_row + 40) + 1):
                v = ws.cell(row=row, column=DAY_COL).value
                if isinstance(v, (datetime.datetime, datetime.date)):
                    key = (v.year, v.month)
                    counts[key] = counts.get(key, 0) + 1
        if not counts:
            return None
        return max(counts.items(), key=lambda kv: kv[1])[0]

    def headers(self, sheet_name: str) -> dict[int, str | None]:
        """{列番号: 見出しの氏名}。空欄・「非常勤」等は None（＝セル内の名前を使う列）。"""
        ws = self._wb[sheet_name]
        result: dict[int, str | None] = {}
        for col in range(FIRST_PERSON_COL, ws.max_column + 1):
            raw = ws.cell(row=self.header_row, column=col).value
            name = tu.normalize_person(raw)
            if not name or name in NON_PERSON_HEADERS or any(w in name for w in NON_PERSON_HEADERS):
                result[col] = None
            else:
                result[col] = name
        return result

    def person_columns(self, sheet_name: str) -> dict[str, int]:
        """{見出しの氏名: 列番号}（クロスリファレンス用）。"""
        return {name: col for col, name in self.headers(sheet_name).items() if name}

    def cell(self, sheet_name: str, row: int, col: int):
        return self._wb[sheet_name].cell(row=row, column=col).value

    # ------------------------------------------------------------ セル1つの解釈

    def interpret(self, raw, header_person: str | None, source: str) -> DayEntry | None:
        """セル1つを DayEntry（day は呼び出し側で設定）に変換する。空セルは None。"""
        text = tu.normalize(raw)
        if not text:
            return None

        ranges = tu.parse_time_ranges(text)
        absence = tu.match_absence(text)
        candidates = tu.extract_name_candidates(text)
        gakudo_hits = [c for c in candidates if tu.gakudo_key(c) in self.gakudo_sheets]
        name_hits = [c for c in candidates if c not in gakudo_hits]

        notes: list[str] = []
        warnings: list[str] = []

        # 種別を決める（時刻 → 欠勤系 → 休み → 応援のみ → 不明）
        if ranges:
            kind = "time"
            if len(ranges) > 1:
                notes.append(f"分割シフト {len(ranges)} 件として解釈")
            for start, end in ranges:
                if tu.minutes_between(start, end) == 0:
                    warnings.append(f"開始{start:%H:%M}〜終了{end:%H:%M}が不正（勤務0分）")
        elif absence:
            kind = "absence"
        elif tu.is_holiday_mark(text):
            kind = "holiday"
        elif gakudo_hits:
            kind = "transfer"
        else:
            kind = "unknown"

        destination = gakudo_hits[0] if gakudo_hits else None
        if destination:
            notes.append(f"応援先の記載「{destination}」を検出")

        # 誰の勤務かを決める
        person = header_person
        if person is None:
            if name_hits:
                person = tu.normalize_person(name_hits[0])
                notes.append(f"列見出しが空欄/非常勤のため、セル内の氏名「{person}」として解釈")
            else:
                person = ""
        elif name_hits and kind in ("time", "transfer"):
            # 見出しの先生の列に別の人の名前が書かれている（代打・応援など）ことがある
            embedded = tu.normalize_person(name_hits[0])
            if embedded and embedded != person:
                warnings.append(
                    f"列見出しは「{person}」だがセル内に「{embedded}」の記載あり（誰の勤務か要確認）"
                )

        entry = DayEntry(
            day=0,
            person=person,
            kind=kind,
            shifts=ranges,
            absence=tu.match_absence(text) if kind == "absence" else None,
            destination=destination,
            sources=[source],
            raws=[text],
            notes=notes,
            warnings=warnings,
        )
        if not person:
            entry.warnings.append("氏名が特定できないセル")
            entry.kind = "unknown"
        return entry

    # ------------------------------------------------------------ シート1枚の解釈

    def parse_sheet(self, sheet_name: str) -> list[DayEntry]:
        """①の1シートを走査して DayEntry のリストを返す。

        見出しが空欄の列にも勤務が書かれていることがあるため、D列以降の全列を走査する。
        C列（行事予定等）は学童全体のメモなので対象外。
        """
        ws = self._wb[sheet_name]
        headers = self.headers(sheet_name)
        entries: list[DayEntry] = []
        for day, row in sorted(self.day_rows(sheet_name).items()):
            for col in range(FIRST_PERSON_COL, ws.max_column + 1):
                raw = ws.cell(row=row, column=col).value
                source = f"{sheet_name}!{openpyxl.utils.get_column_letter(col)}{row}"
                entry = self.interpret(raw, headers.get(col), source)
                if entry is None:
                    continue
                entry.day = day
                entries.append(entry)
        return entries

    def parse_all(self, sheet_names: list[str] | None = None) -> list[DayEntry]:
        targets = sheet_names or self.sheet_names
        entries: list[DayEntry] = []
        for name in targets:
            entries.extend(self.parse_sheet(name))
        return entries

    # ------------------------------------------------------------ クロスリファレンス

    def find_person_shifts(self, gakudo_name: str, day: int, person: str) -> tuple[list, str | None]:
        """応援先の学童シートから、指定日・指定の先生の実際の時刻を探す。

        戻り値は (時刻レンジのリスト, 見つけた場所)。見つからなければ ([], None)。
        """
        key = tu.gakudo_key(gakudo_name)
        person_key = tu.normalize_person(person)
        for sheet_name in self.gakudo_sheets.get(key, []):
            rows = self.day_rows(sheet_name)
            if day not in rows:
                continue
            row = rows[day]
            ws = self._wb[sheet_name]
            headers = self.headers(sheet_name)

            # 1) 見出しがその先生の列
            for col, name in headers.items():
                if name != person_key:
                    continue
                ranges = tu.parse_time_ranges(ws.cell(row=row, column=col).value)
                if ranges:
                    loc = f"{sheet_name}!{openpyxl.utils.get_column_letter(col)}{row}"
                    return ranges, loc

            # 2) セル内に氏名が埋め込まれている列（非常勤列など）
            for col in range(FIRST_PERSON_COL, ws.max_column + 1):
                raw = ws.cell(row=row, column=col).value
                text = tu.normalize(raw)
                if not text or person_key not in tu.normalize_person(text):
                    continue
                ranges = tu.parse_time_ranges(text)
                if ranges:
                    loc = f"{sheet_name}!{openpyxl.utils.get_column_letter(col)}{row}"
                    return ranges, loc
        return [], None

    def resolve_transfers(self, entries: list[DayEntry]) -> list[DayEntry]:
        """kind=="transfer"（学童名のみ）のエントリを応援先の時刻で埋める。"""
        for entry in entries:
            if entry.kind != "transfer" or not entry.destination:
                continue
            ranges, loc = self.find_person_shifts(entry.destination, entry.day, entry.person)
            if ranges:
                entry.kind = "time"
                entry.shifts = ranges
                entry.notes.append(f"応援先 {loc} から時刻を取得")
            else:
                entry.warnings.append(
                    f"応援先「{entry.destination}」の{entry.day}日に「{entry.person}」の時刻が見つからない"
                )
        return entries


# ------------------------------------------------------------------ 集約

def merge_entries(entries: list[DayEntry]) -> dict[tuple[str, int], DayEntry]:
    """(氏名, 日) 単位に集約する。

    応援・アルバイトの先生は複数の学童シートに登場しうるので、勤務時間はすべて
    その人の1日分としてまとめる（SPEC.md「名前 → ③ファイルの紐付けルール」3.）。
    """
    merged: dict[tuple[str, int], DayEntry] = {}
    for entry in entries:
        key = (entry.person, entry.day)
        if key not in merged:
            merged[key] = entry
            continue
        base = merged[key]
        for shift in entry.shifts:
            if shift not in base.shifts:
                base.shifts.append(shift)
        base.sources.extend(entry.sources)
        base.raws.extend(entry.raws)
        base.notes.extend(entry.notes)
        base.warnings.extend(entry.warnings)
        if entry.destination and not base.destination:
            base.destination = entry.destination
        if base.shifts:
            if base.kind != "time":
                base.kind = "time"
            if len(base.sources) > 1 and len(base.shifts) > 1:
                base.warnings.append("複数シート/セルに勤務記録あり（重複入力でないか要確認）")
        elif entry.kind == "absence" and base.kind != "absence":
            base.kind = "absence"
            base.absence = entry.absence
        elif base.kind == "unknown" and entry.kind != "unknown":
            base.kind = entry.kind
    for entry in merged.values():
        entry.shifts.sort()
    return merged
