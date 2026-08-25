"""③（学童ごとの給与明細シート）→ ②（全社集計xlsm）の集計・転記。

元のExcel VBAをそのまま移植したロジック（SPEC.md「③→②の集計転記ロジック（確定・検証済み）」）。
「給与」シートは数式なので、data_only=True でキャッシュされた計算結果を読む。
"""

from __future__ import annotations

import os
import shutil
from dataclasses import dataclass, field

import openpyxl

from . import textutil as tu

SALARY_SHEET = "給与"
BLOCK_START_ROW = 8
BLOCK_MAX_ROW = 1000

# ②の書き込み開始行（「磐田」は73行目と判明済み。「浜松」は要確認）
SUMMARY_START_ROWS = {"磐田": 73, "浜松": 73}


class StaleFormulaCacheError(RuntimeError):
    """③の数式に計算結果がキャッシュされていない（Excelで開いて保存し直す必要がある）。"""


@dataclass
class GakudoTotals:
    """③1ファイル（＝1学童）分の事業所合計。"""

    path: str
    aa_salary: float = 0.0      # 給与（処遇改善除く） → ②E列
    bb_kaizen: float = 0.0      # 処遇改善手当           → ②F列
    cc_transit: float = 0.0     # 交通費                 → ②G列
    blocks: list[dict] = field(default_factory=list)  # 先生ごとの内訳（検算用）

    @property
    def filename(self) -> str:
        return os.path.basename(self.path)


def _is_error_value(v) -> bool:
    """openpyxl で読んだ値がExcelのエラー値かどうか（VBAのIsError相当）。"""
    return isinstance(v, str) and v.startswith("#")


def _num(v) -> float:
    if v is None or _is_error_value(v):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def compute_gakudo_totals(kintai_path: str, sheet_name: str = SALARY_SHEET) -> GakudoTotals:
    """③1ファイル分の「給与」シートから aa/bb/cc を求める。

    「給与」シートは先生ごとに約13〜14行の繰り返しブロックで、各ブロックは
    B列に「名前」が現れる行から始まる。i を「名前」の次の行に進めてから相対参照する。

        a = B(i+6)                       基本給
        b = G(i+8)                       処遇改善
        c = B(i+10) + C(i+10)            非税通勤 + 課税通勤
        if E(i+5) == "法定内残":          ※正社員テンプレートのみ存在する項目
            a += E(i+6)   法定内残
            a += I(i+6)   業務手当
            a += H(i+8)   残業手当
            a -= F(i+10)  不労控除
            a += K(i+6)   休業手当
            a = max(a, 0)
    """
    wb = openpyxl.load_workbook(kintai_path, data_only=True,
                                keep_vba=kintai_path.endswith(".xlsm"))
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"{os.path.basename(kintai_path)} に「{sheet_name}」シートがありません")
    ws = wb[sheet_name]

    def cell(r, c):
        return ws.cell(row=r, column=c).value

    totals = GakudoTotals(path=kintai_path)
    saw_block = False
    saw_value = False

    i = BLOCK_START_ROW
    while i <= BLOCK_MAX_ROW:
        if cell(i, 2) == "名前":
            saw_block = True
            label_row = i
            i += 1

            a = _num(cell(i + 6, 2))                                  # 基本給
            b = _num(cell(i + 8, 7))                                  # 処遇改善
            c = _num(cell(i + 10, 2)) + _num(cell(i + 10, 3))         # 非税通勤+課税通勤

            if cell(i + 5, 5) == "法定内残":                          # 正社員テンプレート
                a += _num(cell(i + 6, 5))    # 法定内残
                a += _num(cell(i + 6, 9))    # 業務手当
                a += _num(cell(i + 8, 8))    # 残業手当
                a -= _num(cell(i + 10, 6))   # 不労控除
                a += _num(cell(i + 6, 11))   # 休業手当
                a = max(a, 0.0)

            if a or b or c:
                saw_value = True

            # 氏名は「名前」ラベルの右か下に入っていることが多い（内訳表示用の参考情報）
            name = None
            for candidate in (cell(label_row, 3), cell(i, 2), cell(i, 3)):
                if isinstance(candidate, str) and tu.normalize_person(candidate):
                    name = tu.normalize_person(candidate)
                    break

            totals.aa_salary += a
            totals.bb_kaizen += b
            totals.cc_transit += c
            totals.blocks.append({"row": label_row, "名前": name, "給与": a, "処遇改善": b, "交通費": c})
        i += 1

    wb.close()

    if saw_block and not saw_value:
        raise StaleFormulaCacheError(
            f"{os.path.basename(kintai_path)}: 「{sheet_name}」シートの計算結果が読めません。"
            "Excelで一度開いて保存し直してから再実行してください"
        )
    return totals


def write_to_summary(
    summary_path: str,
    area_sheet: str,
    rows: dict[str, GakudoTotals],
    start_row: int | None = None,
    output_path: str | None = None,
    dry_run: bool = False,
) -> dict[str, str]:
    """②の area_sheet（「磐田」/「浜松」）のD列と学童名が一致する行に aa/bb/cc を書き込む。

    rows は {学童名: GakudoTotals}。戻り値は {学童名: 結果メッセージ}。
    """
    if start_row is None:
        start_row = SUMMARY_START_ROWS.get(area_sheet, 73)

    dest = summary_path
    if output_path and output_path != summary_path and not dry_run:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        shutil.copy2(summary_path, output_path)
        dest = output_path

    wb = openpyxl.load_workbook(dest, data_only=False, keep_vba=dest.endswith(".xlsm"))
    if area_sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"②に「{area_sheet}」シートがありません（{wb.sheetnames}）")
    ws = wb[area_sheet]

    wanted = {tu.gakudo_key(name): (name, totals) for name, totals in rows.items()}
    results = {name: "②に一致する行が見つからない" for name in rows}

    row = start_row
    while ws.cell(row=row, column=1).value not in (None, ""):
        label = ws.cell(row=row, column=4).value
        key = tu.gakudo_key(label) if label else None
        if key in wanted:
            name, totals = wanted[key]
            if not dry_run:
                ws.cell(row=row, column=5).value = totals.aa_salary
                ws.cell(row=row, column=6).value = totals.bb_kaizen
                ws.cell(row=row, column=7).value = totals.cc_transit
            results[name] = (
                f"{area_sheet}!{row}行 に 給与={totals.aa_salary:,.0f} "
                f"処遇改善={totals.bb_kaizen:,.0f} 交通費={totals.cc_transit:,.0f}"
            )
        row += 1

    if not dry_run:
        wb.save(dest)
    wb.close()
    return results
